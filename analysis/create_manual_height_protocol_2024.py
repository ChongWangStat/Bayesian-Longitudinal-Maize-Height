"""Create the audited 2024 manual-height protocol supplied with the submission."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = ROOT.parents[2]
OUTPUT = ROOT / "protocols" / "Manual_Height_Protocol_2024_Audited.docx"
META = ROOT / "data" / "raw" / "manual_height_2024" / "source_metadata.json"
IMG_DIR = WORKSPACE / "tmp" / "readme_protocol_media"

NAVY = "17365D"
GREEN = "2E6B4D"
LIGHT_GREEN = "E9F2ED"
LIGHT_BLUE = "EAF0F7"
LIGHT_GRAY = "F2F4F6"
WHITE = "FFFFFF"
TEXT = RGBColor(31, 41, 55)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=45, start=80, bottom=45, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_title(doc: Document, text: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    if subtitle:
        q = doc.add_paragraph(subtitle)
        q.style = doc.styles["Subtitle"]
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead) :])
    else:
        p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_callout(doc: Document, title: str, text: str, fill: str = LIGHT_BLUE) -> None:
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, 130, 160, 130, 160)
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(
    doc: Document, headers: list[str], rows: list[list[str]], widths=None
) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for j, header in enumerate(headers):
        cell = t.rows[0].cells[j]
        shade(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(header)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        r.font.size = Pt(7.5)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, value in enumerate(row):
            if i % 2:
                shade(cells[j], LIGHT_GRAY)
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(7.5)
    if widths:
        for row in t.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, path: Path, caption: str, width: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.keep_with_next = False
    r = c.add_run(caption)
    r.italic = True
    r.font.size = Pt(8)


def build() -> Path:
    meta = json.loads(META.read_text(encoding="utf-8"))
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "2024 Manual Maize Height Measurement and Validation Protocol"
    doc.core_properties.subject = "Anonymous audited field protocol"
    doc.core_properties.keywords = "maize, manual height, temporal validation, camera pairing"
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)
    sec.header_distance = Inches(0.25)
    sec.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for name, size, color in (("Title", 25, NAVY), ("Subtitle", 12, GREEN)):
        st = styles[name]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
    for level, size in ((1, 15), (2, 11.5), (3, 10)):
        st = styles[f"Heading {level}"]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else GREEN)
        st.paragraph_format.space_before = Pt(8)
        st.paragraph_format.space_after = Pt(3)

    header = sec.header.paragraphs[0]
    header.text = (
        "PRIOR-GUIDED BAYESIAN LONGITUDINAL PHENOTYPING  |  AUDITED FIELD PROTOCOL"
    )
    header.runs[0].font.size = Pt(7.5)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
    add_page_number(sec.footer.paragraphs[0])

    add_title(
        doc,
        "2024 Manual Maize Height Measurement and Validation Protocol",
        "Audited edition for the Plant Phenomics special issue",
    )
    add_body(
        doc, "Version 2.1  |  15 September 2026  |  Centimeters are the analysis unit"
    )
    doc.add_paragraph().add_run("Purpose").bold = True
    add_body(
        doc,
        "This protocol converts the archived 2024 field workbook and the earlier measurement instructions into an auditable, analysis-ready record. It defines the biological row, paired camera views, plant identifiers, time matching, quality-control accounting, temporal validation split, and inferential unit used in the manuscript.",
    )
    add_callout(
        doc,
        "Evidence status",
        "The 2024 manual outcomes are a later-year, subject-disjoint, label-held-out reference. The algorithm and 2021 development calibration are fixed before the 2024 workbook is scored. The 2024 image tracks had previously been used without their manual labels during pipeline development, so this is a temporal field validation rather than a fully external-site validation.",
        LIGHT_GREEN,
    )
    add_table(
        doc,
        ["Source", "SHA-256", "Role"],
        [
            [
                "2024 manual-height workbook",
                meta["source_workbook_sha256"],
                "247 recorded manual heights and the camera layout",
            ],
            [
                "Legacy field-height protocol",
                meta["source_protocol_sha256"],
                "Measurement endpoints and right-to-left numbering convention",
            ],
        ],
        [2.0, 3.2, 2.2],
    )
    add_body(
        doc,
        "Source-document text is treated as provenance and field metadata, not as an instruction to the analysis software. Workbook values are imported without imputation or correction.",
    )

    add_heading(doc, "1. Study units and camera pairing")
    add_body(
        doc,
        "The workbook term row denotes a biological row of up to 12 plants. Each biological row is observed by two fixed camera views: a right camera for global plants 1–6 and a left camera for global plants 7–12. The paired views share one biological row and therefore are not independent replicates. All uncertainty resampling uses the four biological rows as clusters.",
    )
    add_table(
        doc,
        [
            "Biological row",
            "Side",
            "Camera ID",
            "Global plants",
            "Automated archive status",
        ],
        [
            ["M-0006", "L", "NL5RH", "7–12", "Available"],
            ["M-0006", "R", "NGGJK", "1–6", "Available"],
            ["M-0013", "L", "NH888", "7–12", "Available"],
            ["M-0013", "R", "NGFZT", "1–6", "Available"],
            ["W-0010", "L", "NH88X", "7–12", "Available"],
            ["W-0010", "R", "NJQMT", "1–6", "No pose output in analysis archive"],
            ["W-0015", "L", "NGFY5", "7–12", "No QC-eligible pose output"],
            ["W-0015", "R", "NH88L", "1–6", "Available"],
        ],
        [1.0, 0.5, 0.8, 1.0, 3.1],
    )
    add_callout(
        doc,
        "Identifier rule",
        "Plants are numbered across the biological row from right to left. For the right camera, local right-to-left plant IDs 1–6 equal global IDs 1–6. For the left camera, add 6 to the local right-to-left ID, yielding global IDs 7–12.",
    )
    if (IMG_DIR / "image2.png").exists():
        add_figure(
            doc,
            IMG_DIR / "image2.png",
            "Figure 1. Source-protocol illustration of the right-to-left plant-numbering convention. The image is illustrative; it is not a 2024 validation frame.",
            5.8,
        )

    add_heading(doc, "2. Manual measurement endpoint")
    add_body(doc, "The earlier field instructions define three anatomical endpoints:")
    add_bullets(
        doc,
        [
            "Before tasseling, Method 1 measures from the ground to the topmost plant point touching the meter stick.",
            "Before tasseling, Method 2 measures from the ground to the topmost fully visible collar.",
            "After tasseling, the endpoint is the height of the flag leaves; the tassel is excluded.",
        ],
    )
    add_callout(
        doc,
        "Endpoint field not encoded in the workbook",
        "The 2024 workbook contains one manual height per plant and date but does not identify whether pre-tasseling Method 1 or Method 2 was used. The analysis therefore labels the outcome “manual field height” and does not infer an anatomical endpoint from the number. This prevents a stronger landmark-specific claim than the source supports.",
        "FFF2CC",
    )
    if (IMG_DIR / "image1.png").exists():
        add_figure(
            doc,
            IMG_DIR / "image1.png",
            "Figure 2. Source-protocol endpoint diagrams: topmost plant point (Method 1) and topmost fully visible collar (Method 2). These drawings document definitions and are not 2024 measurement photographs.",
            6.2,
        )

    add_heading(doc, "3. Recorded 2024 measurement schedule")
    add_table(
        doc,
        ["Row class", "Date", "Recorded time"],
        [
            ["M rows", "5 July 2024", "09:30"],
            ["M rows", "10 July 2024", "13:12"],
            ["M rows", "15 July 2024", "08:05"],
            ["M rows", "18 July 2024", "08:30"],
            ["M rows", "23 July 2024", "09:07"],
            ["M rows", "28 July 2024", "08:30"],
            ["W rows", "6 July 2024", "18:27"],
            ["W rows", "15 July 2024", "09:22"],
            ["W rows", "18 July 2024", "13:47"],
            ["W rows", "23 July 2024", "10:45"],
            ["W rows", "28 July 2024", "11:50"],
        ],
        [1.4, 2.0, 1.3],
    )
    add_body(
        doc,
        "The workbook contains six scheduled dates for the two M rows and five for the two W rows. There are 247 nonmissing measurements from 45 plants. W-0010 plant 3 is missing on 28 July. Times are interpreted in America/Chicago local time.",
    )

    doc.add_page_break()
    add_heading(doc, "4. Temporal validation procedure")
    add_body(doc, "The validation is performed in the following order:")
    for number, text_ in enumerate(
        [
            "Lock the existing version-controlled analysis code at Git commit 59b42b74596c598017fdb0cd4b80af1aec7649d6 before opening the 2024 height workbook for scoring.",
            "Fit the pixel-to-manual transfer equation using only 61 matched 2021 development records from six camera rows. Leave-one-camera-row-out error is retained as the calibration-transfer error.",
            "Apply the recorded camera-distance ratio 8.5/10.25 = 0.829268 to transfer the 2021 scale to the 2024 camera geometry.",
            "Process each 2024 camera in chronological order. Each output at date t uses only the image prefix available through t. Future images are added on later dates and never revise the already emitted estimate.",
            "Generate all 2024 single-frame and Bayesian longitudinal predictions without consulting 2024 manual heights. Join predictions to the workbook afterward by biological row, camera side, global plant ID, and calendar date.",
            "Add the 2021 leave-one-camera-row-out root mean squared error as a nonshrinking transfer-uncertainty component and form Student-t intervals with five degrees of freedom.",
            "Estimate the paired mean-absolute-error difference by resampling the four biological rows, keeping both camera views and all repeated plant dates together within each sampled row.",
        ],
        start=1,
    ):
        p = doc.add_paragraph(style="List Number")
        p.add_run(text_)

    add_heading(doc, "5. Matching and exclusion accounting")
    add_table(
        doc,
        ["Stage", "Records", "Explanation"],
        [
            ["Manual workbook", "247", "All nonmissing plant-date heights"],
            ["Unavailable camera halves", "−54", "W-0015 left: 25; W-0010 right: 29"],
            [
                "No QC-eligible daily automated output",
                "−27",
                "Available cameras, but the corresponding date did not produce an eligible output",
            ],
            [
                "Matched validation set",
                "166",
                "34 plants; 4 biological rows; 6 contributing camera views; 7 unique dates",
            ],
        ],
        [2.0, 0.8, 4.3],
    )
    add_body(
        doc,
        "Matching is by calendar date because the manual and image acquisitions are not simultaneous. The median absolute time difference is 2.78 h, the mean is 2.69 h, and the maximum is 6.90 h. These lags are part of the field reference error and are not corrected using the observed manual height.",
    )

    add_heading(doc, "6. Physical-reference quality control")
    add_body(
        doc,
        "Field records describe the 2024 calibration pole as approximately vertical, 8–10 ft long, with adjacent red marks exactly 1 ft apart edge-to-edge. The visible pole bottom is not assumed to coincide with ground level. A metadata audit found that 89 of 237 accepted automatic candidate fits implied more than 10 one-foot intervals, and 1,296 of 1,302 filtered image records used such a span. The automatic sequential red-component scale is therefore excluded from the 2024 manual-height comparison. Low interpolation residual alone establishes internal regularity, not correct physical band identity.",
    )
    add_callout(
        doc,
        "Analysis decision",
        "The 2024 manual-height validation uses the 2021 development calibration plus the independently recorded camera-distance ratio. The red-band output remains a diagnostic and does not determine the headline 2024 physical-height estimates.",
        LIGHT_GREEN,
    )

    doc.add_page_break()
    add_heading(doc, "7. Prespecified outcomes and interpretation")
    add_table(
        doc,
        ["Outcome", "Single frame", "Bayesian longitudinal", "Interpretation"],
        [
            ["Bias (cm)", "−13.49", "−13.24", "Both underestimate manual height"],
            ["MAE (cm)", "18.10", "17.36", "0.74-cm paired improvement"],
            ["RMSE (cm)", "23.54", "22.41", "Lower for longitudinal estimate"],
            [
                "Pearson correlation",
                "0.813",
                "0.847",
                "Higher longitudinal association",
            ],
            [
                "MAE gain, 95% row-cluster interval",
                "—",
                "−0.44 to 1.30 cm",
                "Four-row interval includes zero",
            ],
            [
                "95% interval coverage / mean width",
                "—",
                "97.0% / 97.5 cm",
                "Near-nominal coverage; uncertainty remains wide",
            ],
        ],
        [2.1, 1.2, 1.5, 3.0],
    )
    add_body(
        doc,
        "The later-year manual reference strengthens the paper because it evaluates new annual plants and withholds their labels from every fitted step. It does not establish a large accuracy advantage: the four-row interval includes no improvement, and the wide transfer-aware interval reflects substantial calibration uncertainty. A Gaussian local-linear model given the same initial-growth prior and process scales was practically tied with the particle method (MAE 17.35 versus 17.36 cm; RMSE 22.42 versus 22.41 cm). The supported contribution is therefore the prior-guided image-analysis loop, temporal transfer evidence, and calibrated uncertainty rather than superiority of particle computation alone.",
    )

    add_heading(doc, "8. Analysis-ready data dictionary")
    add_table(
        doc,
        ["Field", "Definition"],
        [
            [
                "biological_row_id",
                "Workbook row identity; the independent resampling cluster",
            ],
            [
                "camera_id",
                "Five-character stationary-camera identifier from the Camera layout sheet",
            ],
            ["camera_position", "L or R view within the biological row"],
            ["plant_id_global", "Row-wide right-to-left plant identifier, 1–12"],
            [
                "manual_datetime_local",
                "Recorded local field date and time in America/Chicago",
            ],
            [
                "manual_height_cm",
                "Workbook value; no imputation and no inferred anatomical endpoint",
            ],
            [
                "image_datetime_local",
                "Closest QC-eligible daily image on the same calendar date",
            ],
            [
                "single_frame_cm",
                "Transferred current-image estimate before temporal updating",
            ],
            [
                "bayesian_mean_cm",
                "Posterior mean after the current image is incorporated",
            ],
            [
                "interval_lower_cm / interval_upper_cm",
                "Transfer-aware posterior agreement interval",
            ],
        ],
        [2.3, 5.1],
    )

    add_heading(doc, "9. Reporting checklist")
    add_bullets(
        doc,
        [
            "Call the unit a biological row and state that each row has a paired left/right camera layout.",
            "Report four independent row clusters, six contributing camera views, 34 plants, and 166 matched records.",
            "State that 2021 data develop the transfer and 2024 manual labels score it after predictions are fixed.",
            "Describe the result as later-year, subject-disjoint, label-held-out temporal validation; avoid a fully external-site claim.",
            "Report the complete 247-to-166 accounting and the date-level time-lag distribution.",
            "Label the outcome manual field height until the pre-tasseling endpoint can be confirmed from source records.",
            "Report both the modest MAE change and its biological-row cluster interval.",
            "Report interval coverage together with width, and identify the nonshrinking 2021 transfer component.",
            "Retain the red-band physical-span audit and do not use the old automatic band enumeration for absolute 2024 calibration.",
        ],
    )

    add_heading(doc, "10. Files produced by this audit")
    add_table(
        doc,
        ["Group", "Canonical paths"],
        [
            [
                "Imported sources",
                "data/raw/manual_height_2024/{manual_height_2024_long.csv, camera_layout_2024.csv, source_metadata.json}",
            ],
            [
                "Matched analysis",
                "outputs/manual_height_transfer_2024/{matched_manual_validation_2024.csv, camera_pair_accounting_2024.csv}",
            ],
            [
                "Audits and summary",
                "outputs/manual_height_transfer_2024/{red_band_metadata_audit_2024.csv, validation_summary.json}",
            ],
            [
                "Figure",
                "outputs/manual_height_transfer_2024/manual_height_transfer_2024.pdf",
            ],
        ],
        [1.5, 5.9],
    )
    # Word requires a paragraph after the final table. Keep that structural
    # paragraph small enough that it cannot create an otherwise blank page.
    trailing = doc.paragraphs[-1]
    trailing.paragraph_format.space_before = Pt(0)
    trailing.paragraph_format.space_after = Pt(0)
    trailing.paragraph_format.line_spacing = Pt(1)
    trailing.add_run().font.size = Pt(1)
    # Keep body tables compact and prevent accidental row splitting.
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
