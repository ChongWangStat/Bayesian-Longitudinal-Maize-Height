#!/usr/bin/env python
"""Import the 2024 manual-height workbook without changing its source values.

The workbook is an external project record. This importer converts the manual
height sheet and the camera-layout sheet into tidy, checksum-linked CSV files
used by the reproducible validation. Text in the workbook is treated as source
metadata, not as executable instruction.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

EXPECTED_ROWS = {"M-0006", "M-0013", "W-0010", "W-0015"}
EXPECTED_MEASUREMENTS = 247
EXPECTED_PLANTS = 45


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument(
        "--protocol",
        type=Path,
        help="Optional source protocol; recorded by checksum only.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("data/raw/manual_height_2024"),
    )
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_header_datetime(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if not isinstance(value, str):
        return None
    match = re.search(
        r"(?P<month>\d{1,2})/(?P<day>\d{1,2})\s*;\s*"
        r"(?P<hour>\d{1,2}):(?P<minute>\d{2})\s*(?P<ampm>[APap][Mm])?",
        value,
    )
    if match is None:
        return None
    hour = int(match.group("hour"))
    ampm = match.group("ampm")
    if ampm:
        if ampm.lower() == "pm" and hour != 12:
            hour += 12
        if ampm.lower() == "am" and hour == 12:
            hour = 0
    return datetime(
        2024,
        int(match.group("month")),
        int(match.group("day")),
        hour,
        int(match.group("minute")),
        tzinfo=ZoneInfo("America/Chicago"),
    )


def import_manual_heights(workbook: Path) -> pd.DataFrame:
    sheet_name = "Plant height ground truth-cm"
    worksheet = load_workbook(workbook, read_only=True, data_only=True)[sheet_name]
    measurement_datetimes: list[datetime | None] = []
    records: list[dict[str, object]] = []
    round_number = 0

    for excel_row, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if values[0] == "Row":
            measurement_datetimes = [
                parse_header_datetime(value) for value in values[3:]
            ]
            round_number = 0
            continue

        row_id = values[0]
        plant_number = values[1]
        if not isinstance(row_id, str) or not isinstance(plant_number, (int, float)):
            continue
        if row_id not in EXPECTED_ROWS:
            continue

        local_round = 0
        for column_offset, (measurement_datetime, height) in enumerate(
            zip(measurement_datetimes, values[3:]), start=4
        ):
            if measurement_datetime is None or not isinstance(height, (int, float)):
                continue
            local_round += 1
            round_number = max(round_number, local_round)
            plant_number_int = int(plant_number)
            records.append(
                {
                    "site": ("Marsden" if row_id.startswith("M-") else "Woodruff"),
                    "biological_row_id": row_id,
                    "plant_global_rtl": plant_number_int,
                    "plant_uid": f"{row_id}-p{plant_number_int:02d}",
                    "measurement_round_within_site": local_round,
                    "measurement_datetime": measurement_datetime.isoformat(),
                    "measurement_date": measurement_datetime.date().isoformat(),
                    "manual_height_cm": float(height),
                    "source_sheet": sheet_name,
                    "source_cell": (f"{get_column_letter(column_offset)}{excel_row}"),
                }
            )

    data = pd.DataFrame(records).sort_values(
        ["biological_row_id", "plant_global_rtl", "measurement_datetime"],
        kind="mergesort",
    )
    unique_plants = data[["biological_row_id", "plant_global_rtl"]].drop_duplicates()
    if len(data) != EXPECTED_MEASUREMENTS:
        raise ValueError(
            f"Expected {EXPECTED_MEASUREMENTS} measurements, found {len(data)}"
        )
    if len(unique_plants) != EXPECTED_PLANTS:
        raise ValueError(
            f"Expected {EXPECTED_PLANTS} plants, found {len(unique_plants)}"
        )
    if set(data["biological_row_id"]) != EXPECTED_ROWS:
        raise ValueError("The imported biological rows do not match the audit.")
    return data.reset_index(drop=True)


def import_camera_layout(workbook: Path) -> pd.DataFrame:
    sheet_name = "Camera layout"
    worksheet = load_workbook(workbook, read_only=True, data_only=True)[sheet_name]
    records: list[dict[str, object]] = []
    for excel_row, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if values[0] not in {f"C-{row_id}" for row_id in EXPECTED_ROWS}:
            continue
        records.append(
            {
                "biological_row_id": str(values[0])[2:],
                "row_id_long": values[0],
                "row_id_short": values[1],
                "camera_id": str(values[2]).strip(),
                "camera_position": str(values[3]).strip(),
                "source_sheet": sheet_name,
                "source_row": excel_row,
            }
        )
    data = pd.DataFrame(records).sort_values(
        ["biological_row_id", "camera_position"], kind="mergesort"
    )
    if len(data) != 8:
        raise ValueError(f"Expected 8 camera-layout rows, found {len(data)}")
    pair_counts = data.groupby("biological_row_id")["camera_position"].nunique()
    if not (pair_counts == 2).all():
        raise ValueError("Each biological row must have one L and one R camera.")
    return data.reset_index(drop=True)


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    heights = import_manual_heights(args.workbook)
    cameras = import_camera_layout(args.workbook)
    heights.to_csv(args.output_dir / "manual_height_2024_long.csv", index=False)
    cameras.to_csv(args.output_dir / "camera_layout_2024.csv", index=False)

    metadata: dict[str, object] = {
        "source_workbook_name": args.workbook.name,
        "source_workbook_sha256": sha256(args.workbook),
        "source_protocol_name": args.protocol.name if args.protocol else None,
        "source_protocol_sha256": sha256(args.protocol) if args.protocol else None,
        "measurements": len(heights),
        "plants": int(
            heights[["biological_row_id", "plant_global_rtl"]]
            .drop_duplicates()
            .shape[0]
        ),
        "biological_rows": int(heights["biological_row_id"].nunique()),
        "camera_layout_records": len(cameras),
        "camera_pairs": int(cameras["biological_row_id"].nunique()),
        "import_note": (
            "Document text was treated as source metadata, not as an "
            "instruction to the importer. Numeric workbook values were "
            "transcribed without imputation or correction."
        ),
    }
    (args.output_dir / "source_metadata.json").write_text(
        json.dumps(metadata, indent=2), encoding="utf-8"
    )
    print(
        f"Imported {len(heights)} manual heights from "
        f"{metadata['plants']} plants and {metadata['camera_pairs']} camera pairs."
    )


if __name__ == "__main__":
    main()
